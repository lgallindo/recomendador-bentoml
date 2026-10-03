#!/usr/bin/env python3
"""Gera o deck com arco: cena → sinais → métricas → cesta → pipelines.

Storytelling (razão de ser): a loja já observa demanda e compras juntas;
transforma isso em sugestões ranqueadas na página do produto.

Tom afirmativo: o que o sistema É e FAZ.
Evitar “N de A, um de B” como esqueleto do título.

    just slides
    # ou: uv run python scripts/gerar_slides.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

TINTA = RGBColor(0x1A, 0x1A, 0x1A)
PAPEL = RGBColor(0xF7, 0xF5, 0xF0)
ESCURO = RGBColor(0x1E, 0x2A, 0x2E)
ACENTO = RGBColor(0x0D, 0x7A, 0x6F)
CLARO = RGBColor(0xFF, 0xFF, 0xFF)
APAGADO = RGBColor(0x4A, 0x55, 0x58)
CAIXA = RGBColor(0xE8, 0xF2, 0xF0)
CODIGO_BG = RGBColor(0x24, 0x32, 0x36)

SERIFA = "Georgia"
SANS = "Calibri"
MONO = "Consolas"

L, A = Inches(13.333), Inches(7.5)
SAIDA = Path(__file__).resolve().parent.parent / "slides" / "recomendador-bentoml.pptx"
N_SLIDES = 19


def _run(par, texto, tam, *, cor=TINTA, fonte=SANS, negrito=False):
    r = par.add_run()
    r.text = texto
    r.font.name = fonte
    r.font.size = Pt(tam)
    r.font.color.rgb = cor
    r.font.bold = negrito
    return r


def caixa_texto(slide, x, y, w, h, linhas, tam, *, cor=TINTA, fonte=SANS,
                negrito=False, alinha=PP_ALIGN.LEFT, espaco=8):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(4)
    tf.margin_top = tf.margin_bottom = Pt(2)
    for i, linha in enumerate(linhas if isinstance(linhas, list) else [linhas]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinha
        p.space_after = Pt(espaco)
        _run(p, linha, tam, cor=cor, fonte=fonte, negrito=negrito)
    return tb


def retangulo(slide, x, y, w, h, preenchimento, *, borda=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = preenchimento
    if borda is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = borda
        sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    return sh


def seta_h(slide, x1, x2, y):
    linha = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, x1, y - Emu(20000), x2 - x1 - Inches(0.12), Emu(40000)
    )
    linha.fill.solid()
    linha.fill.fore_color.rgb = ACENTO
    linha.line.fill.background()
    linha.shadow.inherit = False
    tri = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, x2 - Inches(0.22), y - Inches(0.12), Inches(0.28), Inches(0.24)
    )
    tri.fill.solid()
    tri.fill.fore_color.rgb = ACENTO
    tri.line.fill.background()
    tri.shadow.inherit = False


def fundo(slide, cor):
    f = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, L, A)
    f.fill.solid()
    f.fill.fore_color.rgb = cor
    f.line.fill.background()
    spTree = slide.shapes._spTree
    sp = f._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def rodape(slide, n: int, *, cor=APAGADO):
    caixa_texto(
        slide, Inches(0.6), Inches(6.7), Inches(12), Inches(0.35),
        f"{n} / {N_SLIDES}  ·  cena → métricas → cesta → pipelines", 12, cor=cor,
    )


def slide_01_pitch(prs):
    """Cena da vitrine + ranqueador."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, ESCURO)
    caixa_texto(
        s, Inches(0.8), Inches(0.7), Inches(11.7), Inches(1.0),
        "Recomendador no BentoML", 36, cor=CLARO, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.8), Inches(1.9), Inches(11.7), Inches(1.6),
        [
            "A pessoa abre a página de um jarro de Tracunhaém.",
            "Embaixo: outras peças — mesma técnica, associação, demanda.",
            "Ranqueador HTTP: produto da página → lista curta com score e motivo legível.",
        ],
        20, cor=RGBColor(0xD5, 0xE0, 0xDE), espaco=6,
    )
    retangulo(s, Inches(0.8), Inches(4.0), Inches(5.5), Inches(2.0), CODIGO_BG)
    caixa_texto(
        s, Inches(1.0), Inches(4.2), Inches(5.1), Inches(1.7),
        ["ENTRADA", "", "produto_na_pagina · limite · excluir"],
        16, cor=CLARO, espaco=4,
    )
    retangulo(s, Inches(6.9), Inches(4.0), Inches(5.5), Inches(2.0), ACENTO)
    caixa_texto(
        s, Inches(7.1), Inches(4.2), Inches(5.1), Inches(1.7),
        ["SAÍDA", "", "items[] com score, reason, rank"],
        16, cor=CLARO, espaco=4,
    )
    rodape(s, 1, cor=APAGADO)


def slide_02_pacotes(prs):
    """Onde rodar cada sinal — mapa operacional."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.3), Inches(12), Inches(0.5),
        "Pacotes da aula", 28, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.55), Inches(0.85), Inches(12), Inches(0.4),
        "Baseline na raiz; pastas irmãs com artefato e porta próprios.",
        16, cor=APAGADO,
    )
    rows = [
        ("Raiz", ":3000", "recomendador", "técnica + pop + cooc"),
        ("variante-demografica/", ":3001", "recomendador-demo", "perfil + compras por faixa"),
        ("variante-clustering/", ":3002", "recomendador-cluster", "dendrograma + cesta ⊕ cluster"),
    ]
    y0 = Inches(1.45)
    for i, (pasta, porta, art, sinal) in enumerate(rows):
        y = y0 + i * Inches(1.45)
        retangulo(s, Inches(0.55), y, Inches(12.2), Inches(1.25), CAIXA if i % 2 == 0 else ESCURO,
                  borda=ACENTO)
        cor = TINTA if i % 2 == 0 else CLARO
        caixa_texto(
            s, Inches(0.8), y + Inches(0.2), Inches(11.7), Inches(1.0),
            [f"{pasta}   ·   porta {porta}   ·   {art}", sinal],
            17, cor=cor, espaco=5,
        )
    rodape(s, 2)


def slide_03_sinais(prs):
    """Ponte: o que a loja já observa."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.3), Inches(12), Inches(0.5),
        "Sinais que a loja observa", 28, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.55), Inches(1.0), Inches(12), Inches(1.0),
        [
            "Antes de ranquear, a loja já acumula duas leituras do passado.",
            "Elas chegam ao recomendador como matéria-prima da nota.",
        ],
        18, espaco=5,
    )
    retangulo(s, Inches(0.5), Inches(2.3), Inches(5.9), Inches(3.5), CAIXA, borda=ACENTO)
    caixa_texto(
        s, Inches(0.7), Inches(2.5), Inches(5.5), Inches(3.1),
        [
            "Demanda",
            "",
            "Quantas vezes cada peça",
            "já foi pedida.",
            "Mostra o que a vitrine",
            "costuma vender bem.",
        ],
        17, espaco=4,
    )
    retangulo(s, Inches(6.8), Inches(2.3), Inches(5.9), Inches(3.5), ESCURO)
    caixa_texto(
        s, Inches(7.0), Inches(2.5), Inches(5.5), Inches(3.1),
        [
            "Compra conjunta",
            "",
            "Quais peças saíram",
            "na mesma compra.",
            "Mostra pares que",
            "andam juntos.",
        ],
        17, cor=CLARO, espaco=4,
    )
    rodape(s, 3)


def slide_04_metricas(prs):
    """Definições em prosa — sem código, sem schema."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.2), Inches(12), Inches(0.45),
        "Métricas do projeto", 28, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.35),
        "Linguagem comum da ordenação — o que cada número significa na aula.",
        15, cor=APAGADO,
    )
    blocos = [
        ("Popularidade", "Quão pedido é cada produto no histórico da loja. Vira um peso na ordenação."),
        ("Co-ocorrência", "Quão frequentemente dois produtos saem na mesma compra. Reforça pares que andam juntos."),
        ("Score", "Nota única que mistura esses sinais — e, nas variantes, sinais extras — para ranquear candidatos."),
        ("Reason", "Rótulo legível do motivo da sugestão: mesma técnica, complemento por popularidade, ou os rótulos das variantes."),
    ]
    for i, (tit, corpo) in enumerate(blocos):
        y = Inches(1.15) + i * Inches(1.15)
        retangulo(s, Inches(0.5), y, Inches(12.3), Inches(1.0), CAIXA if i % 2 == 0 else ESCURO,
                  borda=ACENTO)
        cor = TINTA if i % 2 == 0 else CLARO
        caixa_texto(
            s, Inches(0.7), y + Inches(0.12), Inches(11.9), Inches(0.8),
            [tit, corpo], 16, cor=cor, negrito=True, espaco=3,
        )
    caixa_texto(
        s, Inches(0.55), Inches(5.9), Inches(12), Inches(0.5),
        "Demográfica acrescenta afinidade de perfil; clustering acrescenta vizinhança de grupo.",
        14, cor=APAGADO,
    )
    rodape(s, 4)


def slide_05_cesta_mercado(prs):
    """Cesta: prosa primeiro, estrutura depois."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Cesta de mercado", 28, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.55), Inches(0.85), Inches(12), Inches(1.3),
        [
            "Numa compra, a cesta é o conjunto de produtos que saíram juntos.",
            "O recomendador usa esse padrão para aprender associação entre itens:",
            "peças que compartilham compras tendem a se reforçar na co-ocorrência.",
        ],
        17, espaco=4,
    )
    retangulo(s, Inches(0.5), Inches(2.4), Inches(12.3), Inches(3.6), CODIGO_BG)
    caixa_texto(
        s, Inches(0.75), Inches(2.55), Inches(11.8), Inches(0.4),
        "Estrutura: lista de cestas; cada cesta = lista de ids de produto",
        15, cor=RGBColor(0xD5, 0xE0, 0xDE),
    )
    caixa_texto(
        s, Inches(0.75), Inches(3.1), Inches(11.8), Inches(2.6),
        [
            '"cestas": [',
            '  ["p01", "p02", "p03"],',
            '  ["p01", "p02"],',
            '  ["p04", "p05", "p06"]',
            "]",
            "",
            "Ordem dentro da lista irrelevante — importa o par na mesma cesta.",
        ],
        17, cor=RGBColor(0xB8, 0xE0, 0xD8), fonte=MONO, espaco=3,
    )
    rodape(s, 5)


def slide_06_metricas_dados(prs):
    """Como os sinais viram números no artefato."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Métricas nos dados", 26, fonte=SERIFA, negrito=True,
    )
    retangulo(s, Inches(0.5), Inches(1.0), Inches(5.9), Inches(5.0), CAIXA, borda=ACENTO)
    caixa_texto(
        s, Inches(0.7), Inches(1.2), Inches(5.5), Inches(4.6),
        [
            "Popularidade (pop)",
            "",
            "Parte da contagem de pedidos",
            "de cada produto",
            "",
            "pop = pedidos / max(pedidos)",
            "",
            "Escala 0…1 no artefato",
            "O serviço usa pop no score",
        ],
        16, espaco=3,
    )
    retangulo(s, Inches(6.8), Inches(1.0), Inches(5.9), Inches(5.0), ESCURO)
    caixa_texto(
        s, Inches(7.0), Inches(1.2), Inches(5.5), Inches(4.6),
        [
            "Co-ocorrência (cooc)",
            "",
            "Parte das cestas de mercado",
            "Conta pares na mesma compra",
            "",
            "Normaliza pelo parceiro",
            "mais frequente de cada item",
            "",
            "Escala 0…1 no artefato",
        ],
        16, cor=CLARO, espaco=3,
    )
    rodape(s, 6)


def slide_07_dois_pipelines(prs):
    """Hero contrast Treino / Inferência."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.2), Inches(12), Inches(0.45),
        "Dois pipelines — treino e inferência", 26, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.4),
        "Dois momentos distintos. Quem treina prepara o artefato; a vitrine consome a API.",
        15, cor=APAGADO,
    )
    retangulo(s, Inches(0.5), Inches(1.25), Inches(6.0), Inches(5.0), ACENTO)
    caixa_texto(
        s, Inches(0.7), Inches(1.4), Inches(5.6), Inches(4.7),
        [
            "TREINO  ·  offline",
            "",
            "MLOps: training / batch",
            "Comando: just treino",
            "Arquivo: treino.py",
            "",
            "Quem: operador (aula, CI, job)",
            "Quando: catálogo ou regra muda",
            "",
            "Lê dados → calcula pop e cooc",
            "→ grava recomendador:…",
        ],
        16, cor=CLARO, espaco=3,
    )
    retangulo(s, Inches(6.8), Inches(1.25), Inches(6.0), Inches(5.0), ESCURO)
    caixa_texto(
        s, Inches(7.0), Inches(1.4), Inches(5.6), Inches(4.7),
        [
            "INFERÊNCIA  ·  online",
            "",
            "MLOps: inference / serving",
            "Comando: just serve",
            "Arquivo: service.py",
            "",
            "Quem: vitrine (cliente HTTP)",
            "Quando: cada página de produto",
            "",
            "Carrega artefato uma vez",
            "→ POST /recomendar → JSON",
        ],
        16, cor=CLARO, espaco=3,
    )
    rodape(s, 7)


def slide_08_so_treino(prs):
    """Sequência só treino."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Só o treino (offline)", 26, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.55), Inches(0.8), Inches(12), Inches(0.4),
        "Ator: operador que prepara o artefato. A vitrine fica de fora deste momento.",
        16, cor=APAGADO,
    )
    passos = [
        ("1", "Operador", "just treino"),
        ("2", "treino.py", "lê produtos + cestas"),
        ("3", "Cálculo", "pop e cooc"),
        ("4", "Model store", "grava recomendador:…"),
        ("5", "Saída", "imprime tag no store"),
    ]
    y, w, h = Inches(1.6), Inches(2.2), Inches(2.6)
    gap = Inches(0.25)
    x0 = Inches(0.45)
    for i, (num, a, b) in enumerate(passos):
        x = x0 + i * (w + gap)
        sh = retangulo(s, x, y, w, h, ACENTO if i % 2 == 0 else ESCURO)
        tf = sh.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, num, 14, cor=CLARO, fonte=SANS, negrito=True)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        _run(p2, a, 15, cor=CLARO, fonte=MONO, negrito=True)
        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        _run(p3, b, 13, cor=CLARO, fonte=SANS)
        if i < len(passos) - 1:
            seta_h(s, x + w, x + w + gap, y + h / 2)
    caixa_texto(
        s, Inches(0.55), Inches(4.7), Inches(12), Inches(1.4),
        [
            "Este passo deixa o artefato pronto para o serviço carregar.",
            "Catálogo novo? Rode just treino de novo — isso é treino.",
        ],
        17, espaco=5,
    )
    rodape(s, 8)


def slide_09_so_inferencia(prs):
    """Sequência só inferência."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Só a inferência (online)", 26, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.55), Inches(0.75), Inches(12), Inches(0.4),
        "Sobe o serviço uma vez; depois a vitrine conversa com a API a cada página.",
        16, cor=APAGADO,
    )
    retangulo(s, Inches(0.5), Inches(1.3), Inches(12.3), Inches(1.5), ESCURO)
    caixa_texto(
        s, Inches(0.7), Inches(1.5), Inches(11.9), Inches(1.2),
        [
            "Uma vez:  just serve  →  carrega recomendador:latest  →  processo HTTP no ar",
        ],
        18, cor=CLARO, espaco=4,
    )
    retangulo(s, Inches(0.5), Inches(3.1), Inches(12.3), Inches(2.9), CAIXA, borda=ACENTO)
    caixa_texto(
        s, Inches(0.7), Inches(3.3), Inches(11.9), Inches(2.5),
        [
            "Loop — cada página de produto",
            "",
            "vitrine  →  POST /recomendar (pagina, limite, excluir)",
            "service.py  →  aplica a regra sobre o artefato já carregado",
            "vitrine  ←  JSON (items, reason, score)",
            "",
            "A cada clique: usa o pickle pronto. O treino fica quieto.",
        ],
        16, espaco=3,
    )
    rodape(s, 9)


def slide_10_artefato_meio(prs):
    """Ponte de dependência via model store."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Artefato no meio — quem depende de quem", 26, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.55), Inches(0.8), Inches(12), Inches(0.4),
        "Depois de separar os pipelines: a seta que os liga é o model store.",
        16, cor=APAGADO,
    )
    pares = [
        ("catalogo.json", "dados/"),
        ("treino.py", "pop + cooc"),
        ("model store", "recomendador:…"),
        ("service.py", "POST /recomendar"),
        ("JSON", "items[]"),
    ]
    y, w, h = Inches(1.8), Inches(2.2), Inches(2.4)
    gap = Inches(0.25)
    x0 = Inches(0.45)
    for i, (a, b) in enumerate(pares):
        x = x0 + i * (w + gap)
        fill = ACENTO if i == 2 else (ESCURO if i % 2 else CAIXA)
        sh = retangulo(s, x, y, w, h, fill, borda=ACENTO if i != 2 and fill == CAIXA else None)
        cor = TINTA if fill == CAIXA else CLARO
        tf = sh.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, a, 14, cor=cor, fonte=MONO, negrito=True)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        _run(p2, b, 13, cor=cor, fonte=SANS)
        if i < len(pares) - 1:
            seta_h(s, x + w, x + w + gap, y + h / 2)
    caixa_texto(
        s, Inches(0.55), Inches(4.7), Inches(12), Inches(1.4),
        [
            "Esquerda do store = treino. Direita do store = inferência.",
            "A vitrine fala com service.py; o JSON de dados alimenta o treino.",
        ],
        17, espaco=5,
    )
    rodape(s, 10)


def slide_11_algoritmo(prs):
    """Cinco passos online."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Algoritmo na inferência (baseline)", 26, fonte=SERIFA, negrito=True,
    )
    passos = [
        "1. Contexto — pagina, limite, excluir",
        "2. Mesma técnica — candidatos",
        "3. Score 0,7×pop + 0,3×cooc — ordena",
        "4. Complemento por popularidade se faltar vaga",
        "5. Resposta JSON (items, reason, rank)",
    ]
    for i, linha in enumerate(passos):
        y = Inches(1.0) + i * Inches(0.95)
        retangulo(s, Inches(0.55), y, Inches(12.2), Inches(0.8),
                  ACENTO if i % 2 == 0 else CAIXA, borda=ACENTO)
        caixa_texto(
            s, Inches(0.8), y + Inches(0.18), Inches(11.7), Inches(0.5),
            linha, 18, cor=CLARO if i % 2 == 0 else TINTA, negrito=True,
        )
    rodape(s, 11)


def slide_12_api(prs):
    """Contrato HTTP."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, ESCURO)
    caixa_texto(
        s, Inches(0.55), Inches(0.3), Inches(12), Inches(0.5),
        "API — POST /recomendar", 28, cor=CLARO, fonte=SERIFA, negrito=True,
    )
    retangulo(s, Inches(0.55), Inches(1.0), Inches(5.9), Inches(4.5), CODIGO_BG)
    caixa_texto(
        s, Inches(0.75), Inches(1.2), Inches(5.5), Inches(4.1),
        [
            "Entrada",
            "",
            "produto_na_pagina  string",
            "limite              int (padrão 4)",
            "excluir             lista de ids",
            "",
            "A própria página fica de fora",
            "da lista de sugestões.",
        ],
        16, cor=CLARO, espaco=4,
    )
    retangulo(s, Inches(6.8), Inches(1.0), Inches(5.9), Inches(4.5), ACENTO)
    caixa_texto(
        s, Inches(7.0), Inches(1.2), Inches(5.5), Inches(4.1),
        [
            "Saída (sucesso)",
            "",
            "items[].product_id, nome,",
            "  tecnica, regiao, score,",
            "  reason, rank",
            "complemento_usado · strategy",
            "modelo (tag do artefato)",
        ],
        16, cor=CLARO, espaco=4,
    )
    caixa_texto(
        s, Inches(0.55), Inches(5.7), Inches(12), Inches(0.6),
        'Id fora do catálogo → {"erro":"produto_inexistente","items":[]}',
        15, cor=RGBColor(0xB8, 0xE0, 0xD8), fonte=MONO,
    )
    rodape(s, 12, cor=APAGADO)


def slide_13_subir(prs):
    """Comandos por pipeline."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Como subir — treino depois serve", 26, fonte=SERIFA, negrito=True,
    )
    retangulo(s, Inches(0.5), Inches(1.0), Inches(6.0), Inches(5.0), ACENTO)
    caixa_texto(
        s, Inches(0.7), Inches(1.2), Inches(5.6), Inches(4.6),
        [
            "Pipeline TREINO",
            "",
            "just treino",
            "",
            "lê dados/catalogo.json",
            "grava no model store",
            "",
            "Rode de novo quando o",
            "JSON ou a regra mudar.",
        ],
        17, cor=CLARO, espaco=4,
    )
    retangulo(s, Inches(6.8), Inches(1.0), Inches(6.0), Inches(5.0), ESCURO)
    caixa_texto(
        s, Inches(7.0), Inches(1.2), Inches(5.6), Inches(4.6),
        [
            "Pipeline INFERÊNCIA",
            "",
            "just serve",
            "http://127.0.0.1:3000",
            "",
            "just curl-exemplo",
            "just curl-complemento",
            "just curl-inexistente",
            "",
            "Swagger na mesma porta.",
        ],
        17, cor=CLARO, espaco=3,
    )
    rodape(s, 13)


def slide_14_exemplo_p01(prs):
    """Resultado esperado do jarro."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Exemplo p01 — jarro de Tracunhaém", 26, fonte=SERIFA, negrito=True,
    )
    retangulo(s, Inches(0.5), Inches(0.95), Inches(5.9), Inches(5.1), CODIGO_BG)
    caixa_texto(
        s, Inches(0.7), Inches(1.15), Inches(5.5), Inches(4.7),
        [
            "just curl-exemplo",
            "",
            '{"produto_na_pagina":"p01",',
            ' "limite":4,"excluir":[]}',
            "",
            "Contas no papel:",
            "material/calculos-trabalhados.md",
        ],
        16, cor=RGBColor(0xB8, 0xE0, 0xD8), fonte=MONO, espaco=4,
    )
    retangulo(s, Inches(6.7), Inches(0.95), Inches(6.0), Inches(5.1), ESCURO)
    caixa_texto(
        s, Inches(6.9), Inches(1.15), Inches(5.6), Inches(4.7),
        [
            "Lista esperada",
            "",
            "1  Prato  · mesma_tecnica",
            "2  Boneca · mesma_tecnica",
            "3  Rendeira · complemento…",
            "4  Xilogravura · complemento…",
            "",
            "complemento_usado: true",
            "(duas outras cerâmicas;",
            " as vagas restantes vêm",
            " da popularidade geral)",
        ],
        15, cor=CLARO, espaco=3,
    )
    rodape(s, 14)


def slide_15_codigo(prs):
    """Nomes no código — offline | online."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, ESCURO)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Nomes no código — offline | online", 24, cor=CLARO, fonte=SERIFA, negrito=True,
    )
    retangulo(s, Inches(0.5), Inches(0.95), Inches(6.0), Inches(5.1), CODIGO_BG)
    caixa_texto(
        s, Inches(0.7), Inches(1.15), Inches(5.6), Inches(4.7),
        [
            "# treino.py  (offline)",
            "",
            "catalogo → pedidos_brutos",
            "  → max_pedidos → pop",
            "cestas → conta → cooc",
            "→ artefato recomendador",
        ],
        16, cor=RGBColor(0xB8, 0xE0, 0xD8), fonte=MONO, espaco=4,
    )
    retangulo(s, Inches(6.8), Inches(0.95), Inches(6.0), Inches(5.1), ACENTO)
    caixa_texto(
        s, Inches(7.0), Inches(1.15), Inches(5.6), Inches(4.7),
        [
            "# service.py  (online)",
            "",
            "recomendar()",
            "candidatos = mesma tecnica",
            "score = 0.7*pop + 0.3*cooc",
            "se faltar → complemento",
        ],
        16, cor=CLARO, fonte=MONO, espaco=4,
    )
    rodape(s, 15, cor=APAGADO)


def slide_16_literatura(prs):
    """Mapa de nomes."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Literatura de recomendação", 26, fonte=SERIFA, negrito=True,
    )
    linhas = [
        ("pop / pedidos_brutos", "popularity baseline"),
        ("cooc / cestas", "co-ocorrência item–item"),
        ("filtro por tecnica", "filtragem por conteúdo"),
        ("0,7×pop + 0,3×cooc", "híbrido ponderado (regra)"),
        ("complemento por pop", "fallback / cobertura"),
        ("reason no JSON", "explicabilidade por regra"),
        ("artefato + POST", "offline compute + online inference"),
    ]
    for i, (esq, dir_) in enumerate(linhas):
        y = Inches(0.9) + i * Inches(0.75)
        retangulo(s, Inches(0.5), y, Inches(5.9), Inches(0.65), CAIXA, borda=ACENTO)
        caixa_texto(
            s, Inches(0.7), y + Inches(0.12), Inches(5.5), Inches(0.45),
            esq, 15, fonte=MONO, negrito=True,
        )
        retangulo(s, Inches(6.7), y, Inches(6.0), Inches(0.65), ESCURO if i % 2 == 0 else ACENTO)
        caixa_texto(
            s, Inches(6.9), y + Inches(0.12), Inches(5.6), Inches(0.45),
            dir_, 15, cor=CLARO,
        )
    rodape(s, 16)


def slide_17_ds_ml(prs):
    """Próximo passo afirmativo."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.25), Inches(12), Inches(0.45),
        "Data Science → Machine Learning", 26, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.55), Inches(0.8), Inches(12), Inches(0.5),
        "Hoje: ciência de dados + serviço. Próximo passo: features, rótulo e métrica.",
        16, cor=APAGADO,
    )
    blocos = [
        (Inches(0.5), "Hoje", "Regra fixa 0,7 / 0,3\npickle de mapas\nreason = nome da regra", ACENTO),
        (Inches(4.7), "Amanhã (ML)", "features + label\ntreino / validação\nMAE · Precision@k", ESCURO),
        (Inches(8.9), "Contrato HTTP", "POST /recomendar\nmesmo I/O\nmuda a origem do score", CAIXA),
    ]
    for x, titulo, sub, fill in blocos:
        cor = CLARO if fill != CAIXA else TINTA
        sh = retangulo(s, x, Inches(1.6), Inches(3.7), Inches(4.2), fill,
                       borda=ACENTO if fill == CAIXA else None)
        tf = sh.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, titulo, 18, cor=cor, fonte=SANS, negrito=True)
        for linha in sub.split("\n"):
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            _run(p2, linha, 14, cor=cor, fonte=SANS)
    rodape(s, 17)


def slide_18_catalogo_variantes(prs):
    """Quando usar cada pasta."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, PAPEL)
    caixa_texto(
        s, Inches(0.55), Inches(0.2), Inches(12), Inches(0.45),
        "Catálogo de variantes", 26, fonte=SERIFA, negrito=True,
    )
    linhas = [
        ("Baseline (raiz :3000)", "mesma técnica · 0,7 pop + 0,3 cooc · complemento"),
        ("Demográfica (:3001)", "cliente_id · demo = preferências + pop na faixa"),
        ("Clustering (:3002)", "ímpar=cesta · par=cluster · dendrograma ward k=4"),
    ]
    for i, (t, d) in enumerate(linhas):
        y = Inches(0.85) + i * Inches(1.35)
        retangulo(s, Inches(0.55), y, Inches(12.2), Inches(1.2),
                  ESCURO if i == 0 else CAIXA, borda=ACENTO)
        cor = CLARO if i == 0 else TINTA
        caixa_texto(
            s, Inches(0.8), y + Inches(0.2), Inches(11.7), Inches(0.85),
            [t, d], 17, cor=cor, negrito=True, espaco=4,
        )
    caixa_texto(
        s, Inches(0.55), Inches(5.2), Inches(12), Inches(1.0),
        [
            "Na aula: regra auditável → raiz · perfil do cliente → demo · grupos + cesta → clustering.",
            "Predição de demanda (prever pedidos) → repo irmão predicao-demanda-bentoml.",
        ],
        15, cor=APAGADO, espaco=4,
    )
    rodape(s, 18)


def slide_19_empacote(prs):
    """Containerize o serviço de inferência."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s, ESCURO)
    caixa_texto(
        s, Inches(0.55), Inches(0.3), Inches(12), Inches(0.5),
        "Empacote — inferência portátil", 26, cor=CLARO, fonte=SERIFA, negrito=True,
    )
    caixa_texto(
        s, Inches(0.55), Inches(0.95), Inches(12), Inches(0.5),
        "bentofile.yaml congela service + deps + modelo. Container = empacote do pipeline online.",
        16, cor=RGBColor(0xD5, 0xE0, 0xDE),
    )
    retangulo(s, Inches(0.5), Inches(1.7), Inches(12.3), Inches(4.3), CODIGO_BG)
    caixa_texto(
        s, Inches(0.75), Inches(1.95), Inches(11.8), Inches(3.9),
        [
            "# depois de just treino",
            "uv run bentoml build",
            "uv run bentoml containerize Recomendador:latest",
            "docker run --rm -p 3000:3000 Recomendador:<tag> serve",
            "",
            "# mesmo POST /recomendar — agora em imagem OCI",
            "just curl-exemplo",
        ],
        18, cor=RGBColor(0xB8, 0xE0, 0xD8), fonte=MONO, espaco=5,
    )
    rodape(s, 19, cor=APAGADO)


def main() -> None:
    prs = Presentation()
    prs.slide_width = L
    prs.slide_height = A
    slide_01_pitch(prs)
    slide_02_pacotes(prs)
    slide_03_sinais(prs)
    slide_04_metricas(prs)
    slide_05_cesta_mercado(prs)
    slide_06_metricas_dados(prs)
    slide_07_dois_pipelines(prs)
    slide_08_so_treino(prs)
    slide_09_so_inferencia(prs)
    slide_10_artefato_meio(prs)
    slide_11_algoritmo(prs)
    slide_12_api(prs)
    slide_13_subir(prs)
    slide_14_exemplo_p01(prs)
    slide_15_codigo(prs)
    slide_16_literatura(prs)
    slide_17_ds_ml(prs)
    slide_18_catalogo_variantes(prs)
    slide_19_empacote(prs)
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    lock = SAIDA.parent / ".~lock.recomendador-bentoml.pptx#"
    if lock.exists():
        try:
            lock.unlink()
        except OSError:
            pass
    prs.save(SAIDA)
    print(f"escrito {SAIDA}  ({SAIDA.stat().st_size} bytes)  slides={N_SLIDES}")


if __name__ == "__main__":
    main()
